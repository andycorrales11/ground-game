"""
Builds the golden-master fixture the projection engine is checked against.

    python -m backend.tests.fixtures.generate_workbook_parity

The workbook ships with a cache of every formula's last computed value. That
cache is an oracle: it is what the engine's arithmetic has to reproduce, and it
was produced by a completely independent implementation (Excel), so agreeing
with it to 1e-9 is real evidence and not a tautology.

This is the **only** place allowed to read computed cells. The engine's own
extractor reads inputs and nothing else -- see `backend/ingest/workbook.py`.

The fixture is committed so the test suite runs without the workbook, which is
gitignored along with the rest of `data/`. Regenerate it when the workbook is
updated; expect the numbers to move, and expect that to be visible in the diff.
"""
import csv
import re
from dataclasses import fields
from pathlib import Path

import openpyxl

from backend import config
from backend.ingest.workbook import (
    POSITION_BLOCKS,
    TEAM_TABS,
    load_workbook_projections,
)
from backend.services.projection_engine import PlayerRates

# Where a stat lands on a team tab. Keys match StatLine field names.
STAT_COLUMNS = {
    'pass_attempts': 'D',
    'completions': 'E',
    'pass_yards': 'F',
    'pass_tds': 'G',
    'interceptions': 'H',
    'rush_attempts': 'I',
    'rush_yards': 'J',
    'rush_tds': 'K',
    'targets': 'L',
    'receptions': 'M',
    'recv_yards': 'N',
    'recv_tds': 'O',
}

# The master sheets, and the column holding the settings-driven `Custom` score.
MASTER_SHEETS = {'QB': 'N', 'RB': 'O', 'WR': 'N', 'TE': 'L'}

# Master column B points back at the team tab cell the row was filled from,
# e.g. '=ARI!A$2'. That reference is the join key -- it is the only thing that
# says which of the workbook's rows are actually scored.
MASTER_REF = re.compile(r"^=(\w+)!A\$?(\d+)$")

FIXTURES = Path(__file__).parent
EXPECTED = FIXTURES / 'workbook_parity.csv'
INPUT_TEAMS = FIXTURES / 'workbook_inputs_teams.csv'
INPUT_PLAYERS = FIXTURES / 'workbook_inputs_players.csv'
WORKBOOK = config.DATA_DIR / '2026-FFB-Projections-0817.xlsx'

FIELDS = ['team', 'row', 'pos', 'name', *STAT_COLUMNS, 'custom_points']


def _cached_custom_points(values) -> dict[tuple[str, int], float]:
    """
    The `Custom` score for every scored player, keyed by (team, team-tab row).

    Rows whose reference points at an empty team-tab cell are skipped: the master
    tables were filled down past the end of several depth charts, so they carry
    rows for players who do not exist.
    """
    formulas = openpyxl.load_workbook(WORKBOOK, data_only=False)
    scores: dict[tuple[str, int], float] = {}

    for sheet_name, points_column in MASTER_SHEETS.items():
        formula_sheet = formulas[sheet_name]
        value_sheet = values[sheet_name]
        column_index = openpyxl.utils.column_index_from_string(points_column)

        for row in range(2, formula_sheet.max_row + 1):
            reference = formula_sheet.cell(row=row, column=2).value
            if not isinstance(reference, str):
                continue
            match = MASTER_REF.match(reference.strip())
            if not match:
                continue

            team, team_row = match.group(1), int(match.group(2))
            name = value_sheet.cell(row=row, column=2).value
            if not isinstance(name, str) or not name.strip():
                continue

            scores[(team, team_row)] = float(
                value_sheet.cell(row=row, column=column_index).value or 0.0
            )

    return scores


def _write(path: Path, fieldnames: list[str], rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w', newline='', encoding='utf-8') as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def _write_inputs() -> None:
    """
    Freezes the engine's *inputs* alongside the expected outputs.

    Without this the parity test could only run on a machine that has the
    workbook, and the workbook is gitignored with the rest of `data/` -- so the
    one test that proves the engine is correct would never run in CI. With both
    halves committed, the test is hermetic: read inputs, compute, compare.

    Floats are written with `repr` throughout, which round-trips exactly. A
    fixture that quietly rounds its own inputs cannot support a 1e-9 assertion.
    """
    inputs = load_workbook_projections(WORKBOOK)

    _write(
        INPUT_TEAMS,
        ['team', 'plays', 'pass_pct', 'rush_ypc'],
        [
            {
                'team': volume.team,
                'plays': repr(volume.plays),
                'pass_pct': repr(volume.pass_pct),
                'rush_ypc': repr(volume.rush_ypc),
            }
            for volume in inputs.volumes.values()
        ],
    )

    rate_fields = [f.name for f in fields(PlayerRates)]
    _write(
        INPUT_PLAYERS,
        rate_fields,
        [
            {
                name: (
                    getattr(player, name)
                    if name in ('team', 'key', 'name', 'pos')
                    else repr(getattr(player, name))
                )
                for name in rate_fields
            }
            for player in inputs.players
        ],
    )


def build() -> int:
    _write_inputs()

    values = openpyxl.load_workbook(WORKBOOK, data_only=True)
    custom_points = _cached_custom_points(values)

    rows = []
    for team in TEAM_TABS:
        sheet = values[team]
        for pos, (first, last) in POSITION_BLOCKS.items():
            for row in range(first, last + 1):
                name = sheet[f'A{row}'].value
                if not isinstance(name, str) or not name.strip():
                    continue
                if (team, row) not in custom_points:
                    # Projected on the team tab but never pulled into a master
                    # table, so the workbook never scores it. The engine does --
                    # there is nothing to check it against, that is all.
                    continue

                record = {
                    'team': team,
                    'row': row,
                    'pos': pos,
                    'name': name.strip(),
                    'custom_points': repr(custom_points[(team, row)]),
                }
                for stat, column in STAT_COLUMNS.items():
                    cell = sheet[f'{column}{row}'].value
                    record[stat] = repr(float(cell) if cell is not None else 0.0)
                rows.append(record)

    _write(EXPECTED, FIELDS, rows)
    return len(rows)


if __name__ == '__main__':
    count = build()
    print(f"Wrote {count} scored players to {EXPECTED}")
    print(f"Wrote engine inputs to {INPUT_TEAMS} and {INPUT_PLAYERS}")
